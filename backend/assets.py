from __future__ import annotations

import base64
import os
from pathlib import Path
from threading import Lock
from uuid import uuid4

import cv2
from fastapi import HTTPException, UploadFile
from openai import APIError, OpenAI
from PIL import Image
from docx import Document
from pypdf import PdfReader

from .models import ProjectAsset, ProjectState
from .storage import ensure_project, load_project, project_dir, safe_filename, save_project
from .time_utils import now_iso


ALLOWED_EXTENSIONS = {
    ".pdf": "pdf",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".webp": "image",
    ".mp4": "video",
    ".mov": "video",
    ".webm": "video",
    ".m4v": "video",
    ".docx": "word",
}
MAX_BYTES = {
    "pdf": 30 * 1024 * 1024,
    "image": 10 * 1024 * 1024,
    "video": 200 * 1024 * 1024,
    "word": 30 * 1024 * 1024,
}
MAX_ASSETS_PER_PROJECT = 20
MAX_PDF_PAGES = 30
MAX_TEXT_CHARS = 12000
PROJECT_ASSET_LOCKS: dict[str, Lock] = {}


def asset_dir(project_id: str) -> Path:
    return project_dir(project_id) / "assets"


def project_asset_lock(project_id: str) -> Lock:
    if project_id not in PROJECT_ASSET_LOCKS:
        PROJECT_ASSET_LOCKS[project_id] = Lock()
    return PROJECT_ASSET_LOCKS[project_id]


def merge_asset_into_project(project_id: str, asset: ProjectAsset, add_if_missing: bool = True) -> ProjectState | None:
    with project_asset_lock(project_id):
        try:
            project = load_project(project_id)
        except FileNotFoundError:
            project = ProjectState(id=project_id)
        if not add_if_missing and not any(item.id == asset.id for item in project.assets):
            return None
        project.assets = [item for item in project.assets if item.id != asset.id] + [asset]
        return save_project(project)


def normalize_asset_type(value: str | None) -> str:
    allowed = {"company_intro", "reference", "content_unit", "brand", "other"}
    return value if value in allowed else "other"


def detect_kind(filename: str) -> tuple[str, str]:
    suffix = Path(filename).suffix.lower()
    kind = ALLOWED_EXTENSIONS.get(suffix)
    if not kind:
        raise HTTPException(status_code=400, detail="Unsupported file type.")
    return suffix, kind


async def upload_asset(project: ProjectState, upload: UploadFile, asset_type: str | None) -> tuple[ProjectState, ProjectAsset]:
    project = ensure_project(project)
    if len(project.assets) >= MAX_ASSETS_PER_PROJECT:
        raise HTTPException(status_code=400, detail=f"最多上传 {MAX_ASSETS_PER_PROJECT} 个素材。")

    original_name = upload.filename or "未命名素材"
    suffix, kind = detect_kind(original_name)
    data = await upload.read()
    if not data:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")
    if len(data) > MAX_BYTES[kind]:
        raise HTTPException(status_code=400, detail=f"{kind} file is too large.")

    folder = asset_dir(project.id)
    folder.mkdir(parents=True, exist_ok=True)
    asset_id = uuid4().hex
    filename = f"{asset_id}_{safe_filename(Path(original_name).stem)}{suffix}"
    path = folder / filename
    path.write_bytes(data)

    asset = ProjectAsset(
        id=asset_id,
        filename=filename,
        original_filename=original_name,
        content_type=upload.content_type or "",
        asset_type=normalize_asset_type(asset_type),
        size=len(data),
        path=str(path),
        status="processing",
    )

    asset.updated_at = now_iso()
    (folder / f"{asset_id}.json").write_text(asset.model_dump_json(indent=2), encoding="utf-8")
    project = merge_asset_into_project(project.id, asset)
    return project, asset


def parse_uploaded_asset(project_id: str, asset_id: str) -> None:
    folder = asset_dir(project_id)
    asset_path = folder / f"{asset_id}.json"
    if not asset_path.exists():
        return

    asset = ProjectAsset.model_validate_json(asset_path.read_text(encoding="utf-8"))
    try:
        _, kind = detect_kind(asset.original_filename)
        path = Path(asset.path)
        if not path.exists():
            raise FileNotFoundError(f"素材文件不存在：{asset.original_filename}")
        parsed = parse_asset(path, kind, asset)
        asset.summary = parsed["summary"]
        asset.extracted_text = parsed.get("extracted_text", "")
        asset.status = "ready"
        asset.error = ""
    except Exception as exc:
        asset.status = "failed"
        asset.error = str(exc)
        asset.summary = f"素材《{asset.original_filename}》已上传，但解析失败：{exc}"

    asset.updated_at = now_iso()
    asset_path.write_text(asset.model_dump_json(indent=2), encoding="utf-8")
    merge_asset_into_project(project_id, asset, add_if_missing=False)


def parse_asset(path: Path, kind: str, asset: ProjectAsset) -> dict[str, str]:
    if kind == "pdf":
        return parse_pdf(path, asset)
    if kind == "image":
        return parse_image(path, asset)
    if kind == "video":
        return parse_video(path, asset)
    if kind == "word":
        return parse_word(path, asset)
    raise ValueError("Unsupported asset kind.")


def parse_pdf(path: Path, asset: ProjectAsset) -> dict[str, str]:
    reader = PdfReader(str(path))
    chunks: list[str] = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            chunks.append(text)
    extracted = "\n\n".join(chunks).strip()
    if not extracted:
        summary = f"PDF《{asset.original_filename}》未抽取到可用文本，可能是扫描版或图片型 PDF。"
        return {"summary": summary, "extracted_text": ""}
    sample = extracted[:MAX_TEXT_CHARS]
    summary = summarize_text(
        f"""请将以下 PDF 文本整理成宣传片策划可用的素材摘要。

文件名：{asset.original_filename}
素材类型：{asset.asset_type}

要求：
- 提炼文档主题和业务信息
- 提炼品牌/产品/服务卖点
- 提炼可用于宣传片的画面或表达线索
- 控制在 500 字以内

PDF 文本：
{sample}
"""
    )
    return {"summary": summary, "extracted_text": sample}


def parse_word(path: Path, asset: ProjectAsset) -> dict[str, str]:
    document = Document(str(path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table_lines: list[str] = []
    for table in document.tables[:10]:
        for row in table.rows[:30]:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_lines.append(" | ".join(cells))
    extracted = "\n".join([*paragraphs, *table_lines]).strip()
    if not extracted:
        return {"summary": f"Word 文档《{asset.original_filename}》未抽取到可用文本。", "extracted_text": ""}
    sample = extracted[:MAX_TEXT_CHARS]
    summary = summarize_text(
        f"""请将以下 Word 文档内容整理成宣传片策划可用的品牌资产摘要。

文件名：{asset.original_filename}
素材类型：{asset.asset_type}

要求：
- 提炼品牌/产品/服务信息
- 提炼可用于宣传片的核心卖点
- 提炼视觉、语气、素材使用线索
- 控制在 500 字以内

Word 文本：
{sample}
"""
    )
    return {"summary": summary, "extracted_text": sample}


def parse_image(path: Path, asset: ProjectAsset) -> dict[str, str]:
    with Image.open(path) as image:
        image.verify()
    summary = summarize_images(
        [path],
        f"""请阅读这张用户上传的图片，并整理成宣传片策划可用的素材摘要。

文件名：{asset.original_filename}
素材类型：{asset.asset_type}

请包含：
- 图片内容描述
- 可见文字/品牌元素
- 画面风格、色彩、构图
- 可用于分镜的视觉线索
- 控制在 400 字以内
""",
    )
    return {"summary": summary, "extracted_text": ""}


def parse_video(path: Path, asset: ProjectAsset) -> dict[str, str]:
    frames = extract_video_frames(path)
    if not frames:
        return {
            "summary": f"视频《{asset.original_filename}》已上传，但未能抽取关键帧。可作为参考样片文件保留。",
            "extracted_text": "",
        }
    summary = summarize_images(
        frames,
        f"""以下图片是用户上传视频的关键帧。请根据关键帧总结视频对宣传片策划的参考价值。

文件名：{asset.original_filename}
素材类型：{asset.asset_type}

请包含：
- 视频可能呈现的场景/人物/产品/空间
- 视觉风格、节奏、镜头语言
- 可借鉴点
- 不要生成逐秒分镜
- 控制在 600 字以内
""",
    )
    for frame in frames:
        frame.unlink(missing_ok=True)
    return {"summary": summary, "extracted_text": ""}


def extract_video_frames(path: Path, count: int = 8) -> list[Path]:
    capture = cv2.VideoCapture(str(path))
    try:
        total = int(capture.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
        if total <= 0:
            return []
        indexes = sorted({0, total - 1, *[int(total * i / (count - 1)) for i in range(count)]})
        frame_paths: list[Path] = []
        frame_dir = path.parent / f"{path.stem}_frames"
        frame_dir.mkdir(parents=True, exist_ok=True)
        for index in indexes[:count]:
            capture.set(cv2.CAP_PROP_POS_FRAMES, index)
            ok, frame = capture.read()
            if not ok:
                continue
            frame_path = frame_dir / f"frame_{index}.jpg"
            cv2.imwrite(str(frame_path), frame)
            frame_paths.append(frame_path)
        return frame_paths
    finally:
        capture.release()


def openai_client(api_key: str | None, base_url: str | None, missing_message: str) -> OpenAI:
    if not api_key:
        raise RuntimeError(missing_message)
    return OpenAI(
        api_key=api_key,
        base_url=base_url,
    )


def llm_client() -> OpenAI:
    return openai_client(
        os.getenv("LLM_API_KEY") or os.getenv("KIMI_API_KEY"),
        os.getenv("LLM_BASE_URL") or os.getenv("KIMI_BASE_URL", "https://open.bigmodel.cn/api/paas/v4/"),
        "LLM_API_KEY is not configured on the server.",
    )


def vision_client() -> OpenAI:
    return openai_client(
        os.getenv("VISION_API_KEY") or os.getenv("LLM_API_KEY") or os.getenv("KIMI_API_KEY"),
        os.getenv("VISION_BASE_URL")
        or os.getenv("LLM_BASE_URL")
        or os.getenv("KIMI_BASE_URL", "https://dashscope.aliyuncs.com/compatible-mode/v1"),
        "VISION_API_KEY is not configured on the server.",
    )


def summarize_text(prompt: str) -> str:
    try:
        response = vision_client().chat.completions.create(
            model=os.getenv("VISION_MODEL") or "qwen3.6-plus",
            messages=[
                {"role": "system", "content": "你是专业宣传片策划素材整理助手。"},
                {"role": "user", "content": prompt},
            ],
        )
    except APIError as exc:
        raise RuntimeError(f"素材摘要生成失败：{exc.message}") from exc
    content = response.choices[0].message.content
    return (content or "").strip() or "素材已上传，但模型未返回摘要。"


def summarize_images(paths: list[Path], prompt: str) -> str:
    content: list[dict[str, object]] = [{"type": "text", "text": prompt}]
    for path in paths:
        mime = "image/png" if path.suffix.lower() == ".png" else "image/jpeg"
        encoded = base64.b64encode(path.read_bytes()).decode("ascii")
        content.append({"type": "image_url", "image_url": {"url": f"data:{mime};base64,{encoded}"}})
    try:
        response = vision_client().chat.completions.create(
            model=os.getenv("VISION_MODEL") or "qwen3.6-plus",
            messages=[
                {"role": "system", "content": "你是专业宣传片策划素材视觉分析助手。"},
                {"role": "user", "content": content},
            ],
        )
    except APIError as exc:
        raise RuntimeError(f"视觉素材摘要生成失败：{exc.message}") from exc
    content_text = response.choices[0].message.content
    return (content_text or "").strip() or "素材已上传，但视觉模型未返回摘要。"


def delete_asset(project: ProjectState, asset_id: str) -> ProjectState:
    project = ensure_project(project)
    with project_asset_lock(project.id):
        asset = next((item for item in project.assets if item.id == asset_id), None)
        if not asset:
            raise FileNotFoundError(f"Asset not found: {asset_id}")
        folder = asset_dir(project.id).resolve()
        path = Path(asset.path).resolve()
        if folder not in path.parents:
            raise ValueError("Unsafe asset path.")
        path.unlink(missing_ok=True)
        (folder / f"{asset_id}.json").unlink(missing_ok=True)
        frames = folder / f"{path.stem}_frames"
        if frames.exists():
            for child in frames.glob("*"):
                child.unlink(missing_ok=True)
            frames.rmdir()
        project.assets = [item for item in project.assets if item.id != asset_id]
        save_project(project)
        return project


def asset_download_path(project: ProjectState, asset_id: str) -> Path:
    asset = next((item for item in project.assets if item.id == asset_id), None)
    if not asset:
        raise FileNotFoundError(f"Asset not found: {asset_id}")
    folder = asset_dir(project.id).resolve()
    path = Path(asset.path).resolve()
    if folder not in path.parents or not path.exists():
        raise FileNotFoundError(f"Asset file not found: {asset_id}")
    return path
